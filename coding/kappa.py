"""Agreement between the first coding (methods_table.csv) and a second coder's sheet.
Usage: python kappa.py second_coding_<name>.xlsx methods_table.csv"""
import sys, csv
from collections import Counter
from openpyxl import load_workbook
ws = load_workbook(sys.argv[1], data_only=True)['Coding']
first = {f"M{int(r['row']):03d}": r for r in csv.DictReader(open(sys.argv[2], encoding='utf8'))}
pairs = []
for row in ws.iter_rows(min_row=6, values_only=True):
    if not row[1] or not row[5]: continue
    a = first[row[1]]['status']
    b = str(row[5]).strip().upper().replace('N/A', 'N/A')
    a = {'SUPERVISED': 'SUP', 'UNRESOLVED': 'UNRES'}.get(a, a)
    pairs.append((row[1], a, b))
def kappa(ps):
    n = len(ps); po = sum(a == b for a, b in ps) / n
    ca, cb = Counter(a for a, _ in ps), Counter(b for _, b in ps)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    return po, (po - pe) / (1 - pe) if pe < 1 else float('nan')
po, k = kappa([(a, b) for _, a, b in pairs])
print(f'n = {len(pairs)} methods; observed agreement = {po:.2f}; Cohen kappa (status) = {k:.2f}')
print('Disagreements:')
for m, a, b in pairs:
    if a != b: print(f'  {m}: first = {a}, second = {b}')
